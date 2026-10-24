import unittest
import tempfile
from pathlib import Path
from citation_checker import extract_cites, get_ref, match_citations, audit_docx, norm

class CitationTests(unittest.TestCase):
    def keys(self,text):
        return [(c['lead'],c['year']) for c in extract_cites(text,1)]
    def test_formats(self):
        cases=[('(Smith, 2020)',[('smith','2020')]),('Smith (2020)',[('smith','2020')]),
          ('(Smith and Jones, 2020; Brown & White 2021)',[('smith','2020'),('brown','2021')]),
          ('(Smith et al., 2020)',[('smith','2020')]),('(Smith, 2020a, 2020b)',[('smith','2020a'),('smith','2020b')]),
          ('Smith et al. (2020, 2021)',[('smith','2020'),('smith','2021')]),
          ('Smith, Jones, and Brown (2020)',[('smith','2020')]),
          ('(World Health Organization 2020)',[('world health organization','2020')]),
          ('De Jong (2013)',[('de jong','2013')]),('(García-Márquez, 1999)',[('garcia-marquez','1999')]),
          ('(van Merriënboer,\n2020)',[('van merrienboer','2020')]),
          ('(Luo (Alex), 2024)',[('luo','2024')])]
        for text,expected in cases:
            with self.subTest(text=text): self.assertEqual(self.keys(text),expected)
    def test_repeated(self):
        self.assertEqual(len(self.keys('(Smith 2020) and (Smith 2020)')),2)
    def test_non_citations(self):
        for text in ['In 2020 things changed.','(2020)','(In 2020)','(n = 2020)','cost $2020']:
            self.assertEqual(self.keys(text),[],text)
    def test_reference_styles(self):
        for text,aa in [('Smith, J., & Jones, B. (2020). Title.', ['smith','jones']),
                         ('Smith, J., and B. Jones. 2020. Title.', ['smith','jones']),
                         ('World Health Organization. (2020). Title.', ['world health organization'])]:
            self.assertEqual(get_ref(text,1)['authors'],aa)
    def test_ambiguity_not_cited(self):
        rs=[get_ref('Smith, J. (2020). One.',1),get_ref('Smith, K. (2020). Two.',2)]
        cs=extract_cites('(Smith 2020)',1);match_citations(cs,rs)
        self.assertEqual(cs[0]['status'],'Ambiguous')
        self.assertTrue(all(r['status']=='NEEDS_REVIEW' for r in rs))
    def test_coauthor_mismatch(self):
        rs=[get_ref('Smith, J., & Jones, B. (2020). Title.',1)]
        cs=extract_cites('(Smith and Brown 2020)',1);match_citations(cs,rs)
        self.assertEqual(cs[0]['status'],'NEEDS_REVIEW')
    def test_year_and_spelling(self):
        for citation in ['(Smith 2021)','(Smith 2020a)','(Smyth 2020)']:
            rs=[get_ref('Smith, J. (2020). Title.',1)]
            cs=extract_cites(citation,1);match_citations(cs,rs)
            self.assertEqual(cs[0]['status'],'NEEDS_REVIEW')
    def test_document_runs_tables_continuation(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'fixture.docx';d=Document()
            x=d.add_paragraph();x.add_run('(Smith');x.add_run(', 2020)')
            d.add_table(rows=1,cols=1).cell(0,0).text='Jones (2021)'
            d.add_paragraph('References');d.add_paragraph('Smith, J. (2020). Title.')
            d.add_paragraph('https://doi.org/10.1234/test')
            d.add_paragraph('Jones, B. (2021). Other.');d.save(p)
            stats,cs,rs,_=audit_docx(p)
            self.assertEqual((stats['matched'],len(rs)),(2,2))
            self.assertIn('10.1234/test',rs[0]['ref_text'])
    def test_duplicate(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'d.docx';d=Document();d.add_paragraph('References')
            for _ in range(2):d.add_paragraph('Smith, J. (2020). Title. https://doi.org/10.1234/a')
            d.save(p);self.assertEqual(len(audit_docx(p)[3]),1)

if __name__=='__main__': unittest.main()
